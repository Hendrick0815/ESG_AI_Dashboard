"""產生論文 Word 檔的排版工具（python-docx）。中文 標楷體、英數 Times New Roman；A4、表格深藍表頭白字＋隔列淺底。"""
from __future__ import annotations

import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

KAI, TNR = "標楷體", "Times New Roman"
LAT = re.compile(r"[\x20-\x7E−]+")
TEXT_W = Cm(21.0 - 2 * 2.54)
SHOT_DPR = 3
HEAD_BG, STRIPE_BG = "1F3864", "EEF1F6"   # 表頭深藍白字、隔列淺底（與網站截圖一致）
SHOT_CSS_W = 760          # 截圖時表格容器寬度（px）；對應文件全寬 15.9 cm


def table_key(caption):
    """「表 4-1　績效比較」→「4-1」。"""
    m = re.match(r"表\s*([0-9A-Z]+-[0-9]+[a-z]?)", caption or "")
    return m.group(1) if m else None


def _fonts(run, name):
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), name)


def add_text(p, text, bold=False, size=12, italic=False, color=None):
    """依字元種類切成英數段（Times New Roman）與中文段（標楷體），各自設定字型。"""
    pos = 0
    segs = []
    for m in LAT.finditer(text):
        if m.start() > pos:
            segs.append((text[pos:m.start()], KAI))
        seg = m.group(0)
        segs.append((seg, TNR if re.search(r"[A-Za-z0-9]", seg) else KAI))
        pos = m.end()
    if pos < len(text):
        segs.append((text[pos:], KAI))
    for s, f in segs:
        r = p.add_run(s)
        _fonts(r, f)
        r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def rich(p, text, size=12, base_bold=False):
    """支援 **粗體** 標記。"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            add_text(p, part[2:-2], bold=True, size=size)
        elif part:
            add_text(p, part, bold=base_bold, size=size)
    return p


def spacing(p, before=0, after=6, line=1.5, keep_next=False, align=None, first=None, left=None, hanging=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    if first is not None:
        pf.first_line_indent = first
    if left is not None:
        pf.left_indent = left
    if hanging is not None:
        pf.left_indent = (left or 0) + hanging; pf.first_line_indent = -hanging
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def borders(table, sz=4, color="000000", none=False):
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil" if none else "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def cell_margins(table, top=50, bottom=50, left=80, right=80):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def row_props(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit"); trPr.append(c)
    if header:
        h = OxmlElement("w:tblHeader"); trPr.append(h)


class Report:
    def __init__(self, shots=None):
        self.doc = Document()
        self._cap, self._just_table = None, False
        self.web_tables = []            # 本次產生的所有表格（給網站頁面用）
        self.shots = shots or {}        # 表號 → 網站截圖路徑
        st = self.doc.styles["Normal"]
        st.font.name = TNR; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), KAI)
        for lvl, sz in ((1, 18), (2, 15), (3, 13)):
            hs = self.doc.styles[f"Heading {lvl}"]
            hs.font.name = TNR; hs.font.size = Pt(sz); hs.font.bold = True
            hs.font.color.rgb = RGBColor(0, 0, 0)
            rpr = hs.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
            for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(k), KAI if k == "w:eastAsia" else TNR)
            if rf.getparent() is None:
                rpr.insert(0, rf)
            for t in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(t)) is not None:
                    del rf.attrib[qn(t)]
        sec = self.doc.sections[0]
        self._page(sec)
        self.headings = []        # (level, text)
        self.first_h1_in_section = True

    def _page(self, sec):
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.54)
        sec.footer_distance = Cm(1.27)

    def new_section(self, fmt=None, start=1, footer=True):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._page(sec)
        sectPr = sec._sectPr
        pg = OxmlElement("w:pgNumType")
        if fmt:
            pg.set(qn("w:fmt"), fmt)
        pg.set(qn("w:start"), str(start)); sectPr.append(pg)
        sec.footer.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in list(fp.runs):
            r._element.getparent().remove(r._element)
        if footer:
            fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
            r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for k in ("w:ascii", "w:hAnsi", "w:eastAsia"):
                rf.set(qn(k), TNR)
            sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22")
            rpr.append(rf); rpr.append(sz); r.append(rpr)
            t = OxmlElement("w:t"); t.text = "1"; r.append(t); fld.append(r)
            fp._p.append(fld)
        self.first_h1_in_section = True
        return sec

    # ---------------- 內容
    def h(self, level, text):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        if level == 1:
            spacing(p, 0, 18, 1.3, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            if not self.first_h1_in_section:
                p.paragraph_format.page_break_before = True
            self.first_h1_in_section = False
        elif level == 2:
            spacing(p, 18, 8, 1.3, keep_next=True)
        else:
            spacing(p, 12, 6, 1.3, keep_next=True)
        add_text(p, text, bold=True, size={1: 18, 2: 15, 3: 13}[level])
        self.headings.append((level, text))
        return p

    def p(self, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, after=6):
        p = self.doc.add_paragraph()
        spacing(p, 0, after, 1.6, align=align, first=Pt(24) if indent else None)
        rich(p, text, size=size)
        return p

    def bullets(self, items, size=12):
        for it in items:
            p = self.doc.add_paragraph()
            spacing(p, 0, 3, 1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left=Pt(24), hanging=Pt(14))
            add_text(p, "• ", size=size)
            rich(p, it, size=size)

    def caption(self, text):
        self._cap = text
        p = self.doc.add_paragraph()
        spacing(p, 10, 4, 1.3, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, text, bold=True, size=12)

    def note(self, text):
        if self.web_tables and self.web_tables[-1].get("note") is None and self._just_table:
            self.web_tables[-1]["note"] = text
        self._just_table = False
        p = self.doc.add_paragraph()
        spacing(p, 3, 10, 1.25, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        rich(p, text, size=10)

    def table(self, header, rows, widths, align=None, size=10.5, header_size=None):
        """widths：各欄比例；align：'l'/'c'/'r' 字串（每欄一字）。
        有對應的網站截圖（shots[表號]）時改放截圖；表格內容另記在 web_tables，供網站「論文正式版本」頁顯示。"""
        key = table_key(self._cap)
        if key and not any(t["key"] == key for t in self.web_tables):
            self.web_tables.append({"key": key, "caption": self._cap, "header": [str(h) for h in header],
                                    "rows": [[str(v) for v in r] for r in rows], "widths": list(widths),
                                    "align": align or "l" * len(header), "note": None})
        self._just_table = True
        if key in self.shots:
            self.image(self.shots[key], width_cm=self.shot_width(self.shots[key]))
            return None
        n = len(header)
        align = align or "l" * n
        tot = sum(widths)
        cw = [int(TEXT_W * w / tot) for w in widths]
        t = self.doc.add_table(rows=1 + len(rows), cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        borders(t); cell_margins(t)
        A = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
        for i, row in enumerate([header] + rows):
            row_props(t.rows[i], header=(i == 0))
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.width = Emu(cw[j])
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                par = c.paragraphs[0]
                spacing(par, 0, 0, 1.15, align=A["c"] if i == 0 else A[align[j]])
                add_text(par, str(val), bold=(i == 0), size=header_size or size, color="FFFFFF" if i == 0 else None)
                if i == 0:
                    shade(c, HEAD_BG)
                elif i % 2 == 0:
                    shade(c, STRIPE_BG)
        # 欄寬（tblGrid）
        grid = t._tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(cw[j] / 635)))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    @staticmethod
    def shot_width(path):
        """截圖以 SHOT_DPR 倍解析度拍攝；依實際寬度換算成文件寬度（容器全寬 = 15.9 cm）。"""
        from PIL import Image
        w = Image.open(path).size[0] / SHOT_DPR
        return round(min(15.9, 15.9 * w / SHOT_CSS_W), 2)

    def image(self, path, width_cm=15.5):
        p = self.doc.add_paragraph()
        spacing(p, 6, 2, 1.0, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(str(path), width=Cm(width_cm))

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self, pages):
        """pages：{index: 頁碼字串}；headings 只列一、二級（不含致謝、目錄）。"""
        items = [(l, t) for l, t in self.toc_items]
        for i, (lvl, text) in enumerate(items):
            p = self.doc.add_paragraph()
            spacing(p, 6 if lvl == 1 else 0, 0, 1.5, left=None if lvl == 1 else Pt(24))
            p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            add_text(p, text, bold=(lvl == 1), size=13 if lvl == 1 else 12)
            add_text(p, "\t" + str(pages.get(i, "")), bold=(lvl == 1), size=13 if lvl == 1 else 12)

    def save(self, path):
        self.doc.save(str(path))
